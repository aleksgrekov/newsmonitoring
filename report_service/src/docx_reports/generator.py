import asyncio
from datetime import datetime
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.document import Document as DocumentType
from src.docx_reports.formatter import NewsFormatter
from src.logger.logger_config import configure_logging
from src.models.news_model import News

logger = configure_logging(__name__)


class ReportGenerator:
    """Генерирует отчет по новостям."""

    def __init__(self, formatter: NewsFormatter) -> None:
        self.formatter = formatter

    async def create_report(self, news_data: Sequence[News]) -> str:
        """
        Создает и сохраняет отчет в файл.

        Args:
            news_data: Список новостей для отчета.

        Returns:
            Путь к сохраненному файлу.
        """
        # Формируем название файла: "YYYY-MM-DD_news.docx"
        today_date = datetime.now().strftime("%Y-%m-%d")
        report_filename = f"{today_date}_news.docx"

        report_folder = Path(__file__).resolve().parent.parent.parent / "reports"
        report_folder.mkdir(parents=True, exist_ok=True)

        # Полный путь к файлу
        path_for_save = str(report_folder / report_filename)

        # Генерируем отчет
        doc = self._generate_report(news_data)
        await asyncio.to_thread(doc.save, path_for_save)

        return str(path_for_save)

    def _generate_report(self, news_data: Sequence[News]) -> DocumentType:
        """Создает отчет по переданным новостям."""
        doc = self._create_document()

        for news in news_data:
            self._add_news_to_document(doc, news)

        return doc

    @staticmethod
    def _create_document() -> DocumentType:
        """Создает новый документ с заголовком отчета."""
        doc = Document()
        doc.add_heading("Отчет по новостям", level=1)
        return doc

    def _add_news_to_document(self, doc: DocumentType, news: News) -> None:
        """Добавляет новость в документ."""
        doc.add_heading(news.title, level=2)

        if len(news.content) < 500:
            news_content = news.content
        else:
            news_content = news.content[:500] + "..."

        doc.add_paragraph(news_content)

        doc.add_paragraph()

        if news.pub_date:
            paragraph = doc.add_paragraph()
            self.formatter.add_bold_run(paragraph, "Дата публикации: ")

            formatted_date = news.pub_date.strftime("%d.%m.%Y %H:%M")
            paragraph.add_run(formatted_date)

        self.formatter.format_translation(news, doc)

        self.formatter.format_analysis(news, doc)
        doc.add_paragraph("—" * 39)
