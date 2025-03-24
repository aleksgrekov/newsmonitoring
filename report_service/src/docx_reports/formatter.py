from docx.document import Document as DocumentType
from docx.text.paragraph import Paragraph
from src.models.news_model import News


class NewsFormatter:
    """Форматирует новости для вставки в отчет."""

    @staticmethod
    def add_bold_run(paragraph: Paragraph, text: str) -> None:
        """Добавляет жирный текст в параграф."""
        run = paragraph.add_run(text)
        run.bold = True

    def format_translation(self, news: News, doc: DocumentType) -> None:
        """Добавляет перевод новости в документ."""
        if news.translations and news.translations[0]:
            translation = news.translations[0]
            if translation.title:
                paragraph = doc.add_paragraph()
                self.add_bold_run(paragraph, "Перевод: ")
                paragraph.add_run(translation.title)

            if translation.content:
                paragraph = doc.add_paragraph()
                paragraph.add_run(translation.content)

    def format_analysis(self, news: News, doc: DocumentType) -> None:
        """Добавляет анализ новости в документ."""
        if news.analysis:
            analysis = news.analysis[0]
            sentiment_text = "нейтральный"
            if analysis.sentiment > 0:
                sentiment_text = "положительный"
            elif analysis.sentiment < 0:
                sentiment_text = "отрицательный"

            paragraph = doc.add_paragraph()
            self.add_bold_run(paragraph, "Эмоциональный тон: ")
            paragraph.add_run(sentiment_text)

            if analysis.keywords:
                paragraph = doc.add_paragraph()
                self.add_bold_run(paragraph, "Ключевые слова: ")
                paragraph.add_run(analysis.keywords)
